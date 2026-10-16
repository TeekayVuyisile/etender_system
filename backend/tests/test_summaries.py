import io
import zipfile

import pymupdf
import pytest
import respx
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.models import Tender
from app.schemas.summary import TenderSummaryContent
from app.services.documents import DocumentError, ExtractedDocument, extract, guess_format
from app.services.excerpts import build_excerpt, score_chunk
from tests.conftest import FakeLLM, InMemorySummaryStore, add_tender, auth_header, register

DOC_URL = "https://www.etenders.gov.za/home/Download?blobName=doc.pdf"
REQUIREMENTS_TEXT = (
    "SCOPE OF WORK: Routine road maintenance on route N1 in Limpopo for 36 months. "
    "Mandatory requirement: bidders must have a CIDB grading of 7CE or higher and be "
    "registered on the CSD. A compulsory briefing session will be held. "
)


def make_pdf(text: str, pages: int = 1) -> bytes:
    pdf = pymupdf.open()
    for _ in range(pages):
        pdf.new_page().insert_textbox(pymupdf.Rect(40, 40, 560, 800), text, fontsize=9)
    return pdf.tobytes()


def make_docx(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Closing date"
    table.rows[0].cells[1].text = "11 November 2026"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_zip(files: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    return buffer.getvalue()


# --- Text extraction ----------------------------------------------------------


def test_extract_pdf_text():
    [document] = extract(make_pdf(REQUIREMENTS_TEXT, pages=2), "bid.pdf", "pdf")

    assert len(document.pages) == 2
    assert "CIDB grading" in document.pages[0]


def test_extract_docx_includes_tables():
    [document] = extract(make_docx(REQUIREMENTS_TEXT), "bid.docx", "docx")

    text = "\n".join(document.pages)
    assert "SCOPE OF WORK" in text
    assert "Closing date | 11 November 2026" in text


def test_extract_zip_reads_supported_files_inside():
    data = make_zip(
        {
            "tender/bid.pdf": make_pdf(REQUIREMENTS_TEXT),
            "tender/spec.docx": make_docx("Specification of works"),
            "tender/prices.xlsx": b"not supported",
        }
    )

    documents = extract(data, "pack.zip", "zip")

    assert sorted(d.title for d in documents) == ["bid.pdf", "spec.docx"]


def test_corrupt_pdf_raises_document_error():
    with pytest.raises(DocumentError):
        extract(b"definitely not a pdf", "bad.pdf", "pdf")


def test_guess_format():
    assert guess_format("Bid Document.PDF") == "pdf"
    assert guess_format("pack.zip") == "zip"
    assert guess_format("prices.xlsx") is None
    assert guess_format("file", declared="docx") == "docx"


# --- Excerpts -----------------------------------------------------------------


def test_short_documents_are_sent_in_full():
    documents = [ExtractedDocument(title="a.pdf", pages=["Short tender text"])]

    excerpt = build_excerpt(documents, budget=1000)

    assert "Short tender text" in excerpt.text
    assert not excerpt.truncated


def test_long_documents_keep_cover_and_relevant_passages():
    filler = "The bidder acknowledges general terms and conditions of contract. " * 12
    pages = (
        ["TENDER 64/2026 Wastewater treatment works. Closing date 13/11/2026 12:00"]
        + [filler] * 40
        + [REQUIREMENTS_TEXT * 2]
        + [filler] * 40
    )
    documents = [ExtractedDocument(title="bid.pdf", pages=pages)]

    excerpt = build_excerpt(documents, budget=3000)

    assert excerpt.truncated
    assert excerpt.chars_used <= 3000
    assert "Closing date 13/11/2026" in excerpt.text  # cover page always kept
    assert "CIDB grading of 7CE" in excerpt.text  # relevant passage selected


def test_boilerplate_forms_score_lower():
    form = "SBD 4 Declaration of interest. Signature of bidder: ________ Witnesses: ________"
    assert score_chunk(form) < score_chunk(REQUIREMENTS_TEXT)


# --- AI output validation -------------------------------------------------------


def test_ai_output_is_cleaned_and_validated():
    content = TenderSummaryContent.model_validate(FakeLLM().response)

    assert content.requirements.cidb[0].class_code == "CE"
    assert content.requirements.cidb[0].min_grade == 7
    assert len(content.requirements.cidb) == 1  # unknown class XX dropped
    assert content.requirements.provinces == ["Limpopo"]  # unknown province dropped
    assert content.requirements.min_years_experience == 5
    assert content.estimated_value == 2_500_000


def test_malformed_ai_output_becomes_empty_fields():
    content = TenderSummaryContent.model_validate(
        {"summary": "x", "scope": "not a list", "requirements": "nope", "estimated_value": "n/a"}
    )

    assert content.scope == []
    assert content.requirements.cidb == []
    assert content.estimated_value is None


# --- API and background job -------------------------------------------------------


@pytest.fixture
def headers(client: TestClient) -> dict[str, str]:
    return auth_header(register(client))


def tender_with_document(db: Session, url: str = DOC_URL, title: str = "bid.pdf") -> Tender:
    return add_tender(
        db,
        title="Routine road maintenance",
        documents=[{"title": title, "url": url, "format": title.rsplit(".", 1)[-1]}],
    )


def test_summary_status_is_none_before_generation(
    client: TestClient, db: Session, headers: dict[str, str]
):
    tender = tender_with_document(db)

    body = client.get(f"/api/tenders/{tender.id}/summary", headers=headers).json()

    assert body["status"] == "none"


@respx.mock
def test_generate_summary_from_documents(
    client: TestClient,
    db: Session,
    headers: dict[str, str],
    llm: FakeLLM,
    summary_store: InMemorySummaryStore,
):
    respx.get(DOC_URL).respond(content=make_pdf(REQUIREMENTS_TEXT))
    tender = tender_with_document(db)

    started = client.post(f"/api/tenders/{tender.id}/summary", headers=headers)
    # TestClient runs background tasks before returning, so the job has finished here.
    result = client.get(f"/api/tenders/{tender.id}/summary", headers=headers).json()

    assert started.status_code == 202
    assert result["status"] == "ready"
    assert result["source"] == "documents"
    assert result["content"]["requirements"]["cidb"] == [{"class_code": "CE", "min_grade": 7}]
    assert result["documents"][0]["title"] == "bid.pdf"
    assert "CIDB grading of 7CE" in llm.prompts[0]
    assert "Routine road maintenance" in llm.prompts[0]  # tender details are included

    db.expire_all()
    assert float(db.get(Tender, tender.id).estimated_value) == 2_500_000


@respx.mock
def test_scanned_documents_fall_back_to_tender_details(
    client: TestClient, db: Session, headers: dict[str, str], llm: FakeLLM
):
    respx.get(DOC_URL).respond(content=make_pdf("", pages=3))  # no text layer
    tender = tender_with_document(db)

    client.post(f"/api/tenders/{tender.id}/summary", headers=headers)
    result = client.get(f"/api/tenders/{tender.id}/summary", headers=headers).json()

    assert result["status"] == "ready"
    assert result["source"] == "details_only"
    assert result["unreadable_documents"] == ["bid.pdf"]
    assert "No readable document text" in llm.prompts[0]


@respx.mock
def test_existing_summary_is_reused_unless_refreshed(
    client: TestClient, db: Session, headers: dict[str, str], llm: FakeLLM
):
    respx.get(DOC_URL).respond(content=make_pdf(REQUIREMENTS_TEXT))
    tender = tender_with_document(db)

    client.post(f"/api/tenders/{tender.id}/summary", headers=headers)
    client.post(f"/api/tenders/{tender.id}/summary", headers=headers)
    assert len(llm.prompts) == 1

    client.post(f"/api/tenders/{tender.id}/summary?refresh=true", headers=headers)
    assert len(llm.prompts) == 2


def test_summary_already_processing_is_not_started_twice(
    client: TestClient,
    db: Session,
    headers: dict[str, str],
    llm: FakeLLM,
    summary_store: InMemorySummaryStore,
):
    tender = tender_with_document(db)
    summary_store.docs[tender.id] = {"status": "processing"}

    body = client.post(f"/api/tenders/{tender.id}/summary", headers=headers).json()

    assert body["status"] == "processing"
    assert llm.prompts == []


def test_llm_failure_is_stored_as_failed(
    client: TestClient, db: Session, headers: dict[str, str], llm: FakeLLM
):
    llm.response = {"summary": ""}  # empty summaries are rejected
    tender = add_tender(db, documents=[])

    client.post(f"/api/tenders/{tender.id}/summary", headers=headers)
    result = client.get(f"/api/tenders/{tender.id}/summary", headers=headers).json()

    assert result["status"] == "failed"
    assert "empty summary" in result["error"]


def test_summary_for_unknown_tender_returns_404(client: TestClient, headers: dict[str, str]):
    assert client.post("/api/tenders/999999/summary", headers=headers).status_code == 404


@respx.mock
def test_workspace_shows_ai_summary(client: TestClient, db: Session, headers: dict[str, str]):
    respx.get(DOC_URL).respond(content=make_pdf(REQUIREMENTS_TEXT))
    tender = tender_with_document(db)
    client.post("/api/workspace", json={"tender_id": tender.id}, headers=headers)

    before = client.get("/api/workspace", headers=headers).json()["items"][0]["ai_summary"]
    client.post(f"/api/tenders/{tender.id}/summary", headers=headers)
    after = client.get("/api/workspace", headers=headers).json()["items"][0]["ai_summary"]

    assert before is None
    assert after == "The municipality needs a contractor to maintain roads in Limpopo."
